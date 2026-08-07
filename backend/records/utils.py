import csv
import io
from django.http import HttpResponse

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Fill, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
    # 螢光筆底色（雙倍薪：國定假日／颱風天）
    HIGHLIGHT_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    SUMMARY_FILL = PatternFill(start_color='E8F0FE', end_color='E8F0FE', fill_type='solid')
except ImportError:
    OPENPYXL_AVAILABLE = False
    HIGHLIGHT_FILL = None
    SUMMARY_FILL = None

def generate_xlsx(records):
    """產生 Excel (.xlsx)，雙倍薪 (rate_multiplier >= 2.0) 的列以螢光底色標示。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '打工薪資報表'

    headers = ['日期', '上班', '下班', '加班開始', '加班結束', '店家', '時薪',
               '薪資倍率', '上班時數', '加班時數', '本薪', '加班費', '津貼', '合計', '備註']
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')

    from collections import defaultdict
    month_totals = defaultdict(int)

    for r in records:
        pay = r.total_pay()
        base_pay = round(float(r.base_hours) * float(r.hourly_wage) * float(r.rate_multiplier))
        overtime_pay = round(float(r.overtime_hours) * float(r.hourly_wage) * float(r.overtime_rate))
        month_totals[r.date.strftime('%Y-%m')] += pay
        ws.append([
            str(r.date), str(r.start_time)[:5] if r.start_time else '',
            str(r.end_time)[:5] if r.end_time else '',
            str(r.overtime_start)[:5] if r.overtime_start else '',
            str(r.overtime_end)[:5] if r.overtime_end else '',
            r.store.name if r.store else '', r.hourly_wage, r.rate_multiplier,
            float(r.base_hours), float(r.overtime_hours), base_pay, overtime_pay,
            r.subsidy, pay, r.notes
        ])
        row_idx = ws.max_row
        # 雙倍薪（國定假日／颱風天）→ 整列螢光底色
        if float(r.rate_multiplier) >= 2.0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = HIGHLIGHT_FILL

    # 每月小計
    ws.append([])
    ws.append(['月份', '當月總計'])
    for col in range(1, 3):
        ws.cell(row=ws.max_row, column=col).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=col).fill = SUMMARY_FILL
    for month_key in sorted(month_totals.keys()):
        y, m = month_key.split('-')
        ws.append([f'{y} 年 {int(m)} 月', month_totals[month_key]])

    # 欄寬
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_csv(records):
    output = io.StringIO()
    output.write('﻿')  # BOM for Excel
    writer = csv.writer(output)
    writer.writerow(['日期', '上班', '下班', '加班開始', '加班結束', '店家', '時薪', '薪資倍率', '上班時數', '加班時數', '本薪', '加班費', '津貼', '合計', '備註'])

    from collections import defaultdict
    month_totals = defaultdict(int)

    for r in records:
        pay = r.total_pay()
        base_pay = round(float(r.base_hours) * float(r.hourly_wage) * float(r.rate_multiplier))
        overtime_pay = round(float(r.overtime_hours) * float(r.hourly_wage) * float(r.overtime_rate))
        month_totals[r.date.strftime('%Y-%m')] += pay
        writer.writerow([
            r.date,
            str(r.start_time)[:5] if r.start_time else '',
            str(r.end_time)[:5] if r.end_time else '',
            str(r.overtime_start)[:5] if r.overtime_start else '',
            str(r.overtime_end)[:5] if r.overtime_end else '',
            r.store.name if r.store else '', r.hourly_wage,
            r.rate_multiplier, r.base_hours, r.overtime_hours,
            base_pay, overtime_pay, r.subsidy, pay, r.notes
        ])

    # Monthly summary rows
    writer.writerow([])
    writer.writerow(['月份', '當月總計'])
    for month_key in sorted(month_totals.keys()):
        y, m = month_key.split('-')
        writer.writerow([f'{y} 年 {int(m)} 月', month_totals[month_key]])

    return output.getvalue()

def generate_backup(records):
    data = []
    for r in records:
        data.append({
            'store': r.store.name if r.store else '',
            'date': str(r.date),
            'start_time': str(r.start_time) if r.start_time else '',
            'end_time': str(r.end_time) if r.end_time else '',
            'hours': str(r.hours),
            'hourly_wage': r.hourly_wage,
            'rate_multiplier': str(r.rate_multiplier),
            'overtime_start': str(r.overtime_start) if r.overtime_start else '',
            'overtime_end': str(r.overtime_end) if r.overtime_end else '',
            'overtime_hours': str(r.overtime_hours),
            'overtime_rate': str(r.overtime_rate),
            'subsidy': r.subsidy,
            'notes': r.notes,
        })
    return data
